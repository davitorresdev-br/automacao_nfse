#!/usr/bin/env python3
"""
organizador_nfse.py
-------------------
Lê os XMLs de NFS-e de uma pasta, extrai os dados fiscais, renomeia (ou copia)
o XML e o PDF correspondente com o padrão "NFSe <numero> - <prestador>" e gera
uma planilha de controle em Excel.

Suporta os layouts ABRASF (maioria dos municípios), Padrao Nacional e
Prefeitura de Sao Paulo, com um modo genérico de fallback.

Uso:
    python organizador_nfse.py                      # simulação (não altera nada)
    python organizador_nfse.py --aplicar            # copia para ./organizadas
    python organizador_nfse.py --aplicar --mover    # renomeia no lugar
    python organizador_nfse.py --debug nota.xml     # lista as tags do XML

Dependência: openpyxl  (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import codecs
import re
import shutil
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Rode:  pip install openpyxl")


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------
# "{*}" é um curinga de namespace (Python 3.8+): funciona com ou sem xmlns.
# Cada campo aceita uma lista de caminhos; vale o primeiro que existir.
# Um caminho iniciado por "@" lê um atributo em vez de um elemento.

LAYOUTS = [
    {
        "nome": "ABRASF",
        "deteccao": ".//{*}InfNfse",
        "no_nota": ".//{*}InfNfse",
        "campos": {
            "numero": ["{*}Numero", "{*}IdentificacaoNfse/{*}Numero"],
            "serie": ["{*}SeriePrestacao", ".//{*}IdentificacaoRps/{*}Serie"],
            "codigo_verificacao": ["{*}CodigoVerificacao"],
            "data_emissao": ["{*}DataEmissao", "{*}DataEmissaoRps"],
            "competencia": ["{*}Competencia"],
            "prestador_nome": [
                ".//{*}PrestadorServico/{*}RazaoSocial",
                ".//{*}Prestador/{*}RazaoSocial",
                ".//{*}DadosPrestador/{*}RazaoSocial",
                ".//{*}PrestadorServico/{*}NomeFantasia",
            ],
            "prestador_doc": [
                ".//{*}PrestadorServico//{*}Cnpj",
                ".//{*}PrestadorServico//{*}Cpf",
                ".//{*}IdentificacaoPrestador//{*}Cnpj",
            ],
            "tomador_nome": [
                ".//{*}TomadorServico/{*}RazaoSocial",
                ".//{*}Tomador/{*}RazaoSocial",
            ],
            "tomador_doc": [
                ".//{*}TomadorServico//{*}Cnpj",
                ".//{*}TomadorServico//{*}Cpf",
            ],
            "valor_servicos": [
                ".//{*}Servico/{*}Valores/{*}ValorServicos",
                ".//{*}ValorServicos",
            ],
            "valor_iss": [".//{*}Valores/{*}ValorIss", ".//{*}ValorIss"],
            "valor_liquido": [".//{*}ValorLiquidoNfse"],
            "iss_retido": [".//{*}Servico/{*}IssRetido", ".//{*}IssRetido"],
            "codigo_servico": [
                ".//{*}Servico/{*}ItemListaServico",
                ".//{*}ItemListaServico",
                ".//{*}CodigoTributacaoMunicipio",
            ],
            "discriminacao": [".//{*}Discriminacao"],
        },
    },
    {
        "nome": "Padrao Nacional",
        "deteccao": ".//{*}infNFSe",
        "no_nota": ".//{*}infNFSe",
        "campos": {
            "numero": ["{*}nNFSe"],
            "serie": [".//{*}serie"],
            "codigo_verificacao": ["@Id"],
            "data_emissao": ["{*}dhProc", ".//{*}dhEmi"],
            "competencia": [".//{*}dCompet"],
            "prestador_nome": [
                ".//{*}emit/{*}xNome",
                ".//{*}prest/{*}xNome",
                ".//{*}emit/{*}xFant",
            ],
            "prestador_doc": [
                ".//{*}emit/{*}CNPJ",
                ".//{*}emit/{*}CPF",
                ".//{*}prest/{*}CNPJ",
            ],
            "tomador_nome": [".//{*}toma/{*}xNome"],
            "tomador_doc": [".//{*}toma/{*}CNPJ", ".//{*}toma/{*}CPF"],
            "valor_servicos": [".//{*}vServPrest/{*}vServ", ".//{*}vServ"],
            "valor_iss": [".//{*}vISSQN"],
            "valor_liquido": [".//{*}vLiq", ".//{*}vLiqNFSe"],
            "iss_retido": [".//{*}tpRetISSQN"],
            "codigo_servico": [".//{*}cTribNac", ".//{*}cServ/{*}cTribNac"],
            "discriminacao": [".//{*}xDescServ"],
        },
    },
    {
        "nome": "Prefeitura de Sao Paulo",
        "deteccao": ".//{*}ChaveNFe",
        "no_nota": ".//{*}NFe",
        "campos": {
            "numero": [".//{*}ChaveNFe/{*}NumeroNFe", ".//{*}NumeroNFe"],
            "serie": [".//{*}SerieRPS"],
            "codigo_verificacao": [".//{*}ChaveNFe/{*}CodigoVerificacao"],
            "data_emissao": [".//{*}DataEmissaoNFe", ".//{*}DataEmissao"],
            "competencia": [".//{*}Competencia"],
            "prestador_nome": [".//{*}RazaoSocialPrestador"],
            "prestador_doc": [
                ".//{*}CPFCNPJPrestador/{*}CNPJ",
                ".//{*}CPFCNPJPrestador/{*}CPF",
            ],
            "tomador_nome": [".//{*}RazaoSocialTomador"],
            "tomador_doc": [
                ".//{*}CPFCNPJTomador/{*}CNPJ",
                ".//{*}CPFCNPJTomador/{*}CPF",
            ],
            "valor_servicos": [".//{*}ValorServicos"],
            "valor_iss": [".//{*}ValorISS"],
            "valor_liquido": [".//{*}ValorLiquido"],
            "iss_retido": [".//{*}ISSRetido"],
            "codigo_servico": [".//{*}CodigoServico"],
            "discriminacao": [".//{*}Discriminacao"],
        },
    },
]

# Último recurso: busca por nome de tag em qualquer lugar do documento.
LAYOUT_GENERICO = {
    "nome": "Generico (verificar!)",
    "deteccao": None,
    "no_nota": ".",
    "campos": {
        "numero": [
            ".//{*}NumeroNfse",
            ".//{*}nNFSe",
            ".//{*}NumeroNFe",
            ".//{*}Numero",
        ],
        "serie": [".//{*}Serie"],
        "codigo_verificacao": [".//{*}CodigoVerificacao"],
        "data_emissao": [".//{*}DataEmissao", ".//{*}dhEmi", ".//{*}dhProc"],
        "competencia": [".//{*}Competencia", ".//{*}dCompet"],
        "prestador_nome": [
            ".//{*}RazaoSocialPrestador",
            ".//{*}PrestadorServico/{*}RazaoSocial",
            ".//{*}emit/{*}xNome",
            ".//{*}RazaoSocial",
        ],
        "prestador_doc": [".//{*}CPFCNPJPrestador/{*}CNPJ", ".//{*}Cnpj"],
        "tomador_nome": [".//{*}RazaoSocialTomador", ".//{*}toma/{*}xNome"],
        "tomador_doc": [".//{*}CPFCNPJTomador/{*}CNPJ"],
        "valor_servicos": [".//{*}ValorServicos", ".//{*}vServ"],
        "valor_iss": [".//{*}ValorIss", ".//{*}ValorISS", ".//{*}vISSQN"],
        "valor_liquido": [".//{*}ValorLiquidoNfse", ".//{*}vLiq"],
        "iss_retido": [".//{*}IssRetido", ".//{*}ISSRetido"],
        "codigo_servico": [".//{*}ItemListaServico", ".//{*}CodigoServico"],
        "discriminacao": [".//{*}Discriminacao", ".//{*}xDescServ"],
    },
}

MAPA_ISS_RETIDO = {"1": "Sim", "2": "Nao", "true": "Sim", "false": "Nao"}


# ---------------------------------------------------------------------------
# Leitura e extração
# ---------------------------------------------------------------------------

def local(tag: str) -> str:
    """Remove o namespace de uma tag: '{http://...}Numero' -> 'Numero'."""
    return tag.rsplit("}", 1)[-1]


def carregar_raiz(caminho: Path):
    """Lê o XML tolerando BOM e XML escapado dentro de envelope SOAP."""
    dados = caminho.read_bytes()
    if dados.startswith(codecs.BOM_UTF8):
        dados = dados[len(codecs.BOM_UTF8):]
    raiz = ET.fromstring(dados)

    if detectar_layout(raiz) is None:
        # alguns webservices devolvem o XML da nota como texto dentro de uma tag
        for el in raiz.iter():
            texto = (el.text or "").strip()
            if texto.startswith("<") and re.search(r"nfse|nfe", texto[:400], re.I):
                try:
                    interno = ET.fromstring(texto.encode("utf-8"))
                except ET.ParseError:
                    continue
                if detectar_layout(interno) is not None:
                    return interno
    return raiz


def detectar_layout(raiz):
    for layout in LAYOUTS:
        if raiz.find(layout["deteccao"]) is not None:
            return layout
        if local(raiz.tag) == local(layout["no_nota"]):
            return layout
    return None


def encontrar_notas(raiz, layout):
    """Um arquivo pode conter várias notas (lote / retorno de consulta)."""
    if layout is LAYOUT_GENERICO:
        return [raiz]
    notas = raiz.findall(layout["no_nota"])
    if not notas and local(raiz.tag) == local(layout["no_nota"]):
        notas = [raiz]
    return notas


def ler(no, caminhos):
    for caminho in caminhos:
        if caminho.startswith("@"):
            valor = no.get(caminho[1:])
            if valor and valor.strip():
                return valor.strip()
            continue
        alvo = no.find(caminho)
        if alvo is not None and alvo.text and alvo.text.strip():
            return alvo.text.strip()
    return None


def para_numero(texto):
    if not texto:
        return None
    texto = texto.strip().replace(" ", "").replace("R$", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return round(float(texto), 2)
    except ValueError:
        return None


def para_data(texto):
    if not texto:
        return None
    texto = re.sub(r"(Z|[+-]\d{2}:?\d{2})$", "", texto.strip()).split(".")[0]
    formatos = (
        "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m", "%Y%m%d",
    )
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def extrair(no, layout):
    campos = layout["campos"]
    dados = {chave: ler(no, caminhos) for chave, caminhos in campos.items()}
    dados["layout"] = layout["nome"]

    for chave in ("valor_servicos", "valor_iss", "valor_liquido"):
        dados[chave] = para_numero(dados.get(chave))
    for chave in ("data_emissao", "competencia"):
        dados[chave] = para_data(dados.get(chave))

    retido = (dados.get("iss_retido") or "").strip().lower()
    dados["iss_retido"] = MAPA_ISS_RETIDO.get(retido, dados.get("iss_retido"))

    if dados["valor_liquido"] is None and dados["valor_servicos"] is not None:
        liquido = dados["valor_servicos"]
        if dados["iss_retido"] == "Sim" and dados["valor_iss"]:
            liquido -= dados["valor_iss"]
        dados["valor_liquido"] = round(liquido, 2)

    for chave in ("prestador_doc", "tomador_doc"):
        if dados.get(chave):
            dados[chave] = re.sub(r"\D", "", dados[chave])
    return dados


# ---------------------------------------------------------------------------
# Nomes de arquivo
# ---------------------------------------------------------------------------

INVALIDOS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def limpar(texto, limite=55):
    texto = INVALIDOS.sub("", texto or "")
    texto = re.sub(r"\s+", " ", texto).strip(" .")
    return texto[:limite].strip(" .")


def montar_nome(dados, zeros=0):
    numero = (dados.get("numero") or "").strip()
    if zeros and numero.isdigit():
        numero = numero.zfill(zeros)
    prestador = limpar(dados.get("prestador_nome"))
    if not numero or not prestador:
        return None
    return limpar(f"NFSe {numero} - {prestador}", limite=120)


def destino(pasta: Path, base: str, extensao: str, reservados: set):
    """Devolve (caminho, ja_existe). O sufixo (2) so resolve colisao na mesma
    execucao; se o arquivo ja esta no destino, a nota e' considerada organizada."""
    nome = f"{base}{extensao}"
    if nome.lower() in reservados:
        contador = 2
        while f"{base} ({contador}){extensao}".lower() in reservados:
            contador += 1
        nome = f"{base} ({contador}){extensao}"
    reservados.add(nome.lower())
    caminho = pasta / nome
    return caminho, caminho.exists()


# ---------------------------------------------------------------------------
# Pareamento com o PDF
# ---------------------------------------------------------------------------

def digitos_do_nome(caminho: Path):
    return {int(bloco) for bloco in re.findall(r"\d+", caminho.stem)}


def achar_pdf(xml: Path, dados, pdfs, usados):
    irmao = xml.with_suffix(".pdf")
    if irmao.exists() and irmao not in usados:
        return irmao

    numero = (dados.get("numero") or "").strip()
    if numero.isdigit():
        alvo = int(numero)
        for pdf in pdfs:
            if pdf not in usados and alvo in digitos_do_nome(pdf):
                return pdf

    codigo = re.sub(r"\W", "", (dados.get("codigo_verificacao") or "")).lower()
    if len(codigo) >= 6:
        for pdf in pdfs:
            if pdf not in usados and codigo in re.sub(r"\W", "", pdf.stem).lower():
                return pdf
    return None


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------

COLUNAS = [
    ("Numero", "numero", 14, None),
    ("Serie", "serie", 8, None),
    ("Emissao", "data_emissao", 12, "DD/MM/YYYY"),
    ("Competencia", "competencia", 12, "MM/YYYY"),
    ("Prestador", "prestador_nome", 38, None),
    ("CNPJ/CPF Prestador", "prestador_doc", 20, "@"),
    ("Tomador", "tomador_nome", 32, None),
    ("CNPJ/CPF Tomador", "tomador_doc", 20, "@"),
    ("Cod. Servico", "codigo_servico", 12, "@"),
    ("Valor Servicos", "valor_servicos", 15, 'R$ #,##0.00'),
    ("ISS", "valor_iss", 12, 'R$ #,##0.00'),
    ("ISS Retido", "iss_retido", 11, None),
    ("Valor Liquido", "valor_liquido", 15, 'R$ #,##0.00'),
    ("Cod. Verif./Chave", "codigo_verificacao", 24, "@"),
    ("Discriminacao", "discriminacao", 45, None),
    ("Layout", "layout", 20, None),
    ("Arquivo XML", "arquivo_xml", 40, None),
    ("Arquivo PDF", "arquivo_pdf", 40, None),
    ("Status", "status", 34, None),
]


def gerar_planilha(linhas, caminho: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"

    fonte = Font(name="Arial", size=10)
    cabecalho_fonte = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="1F4E78")

    for indice, (titulo, _, largura, _) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=indice, value=titulo)
        celula.font = cabecalho_fonte
        celula.fill = cabecalho_fundo
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(indice)].width = largura

    for numero_linha, linha in enumerate(linhas, start=2):
        for indice, (_, chave, _, formato) in enumerate(COLUNAS, start=1):
            valor = linha.get(chave)
            if isinstance(valor, str) and len(valor) > 500:
                valor = valor[:500] + "..."
            celula = ws.cell(row=numero_linha, column=indice, value=valor)
            celula.font = fonte
            if formato:
                celula.number_format = formato

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{len(linhas) + 1}"
    wb.save(caminho)


# ---------------------------------------------------------------------------
# Modo debug
# ---------------------------------------------------------------------------

def modo_debug(caminho: Path):
    raiz = carregar_raiz(caminho)
    layout = detectar_layout(raiz) or LAYOUT_GENERICO
    print(f"Arquivo : {caminho.name}")
    print(f"Raiz    : {raiz.tag}")
    print(f"Layout  : {layout['nome']}")
    print(f"Notas   : {len(encontrar_notas(raiz, layout))}\n")
    for elemento in raiz.iter():
        texto = (elemento.text or "").strip()
        if texto:
            print(f"  {local(elemento.tag):<32} -> {texto[:70]}")


# ---------------------------------------------------------------------------
# Principal
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Organiza XMLs/PDFs de NFS-e e gera planilha de controle.",
    )
    parser.add_argument("pasta", nargs="?", default=".", help="pasta com os arquivos")
    parser.add_argument("--saida", default="organizadas", help="pasta de destino")
    parser.add_argument("--planilha", default="Controle_NFSe.xlsx")
    parser.add_argument("--aplicar", action="store_true",
                        help="executa de verdade (sem isso, apenas simula)")
    parser.add_argument("--mover", action="store_true",
                        help="renomeia no lugar em vez de copiar para a pasta de saida")
    parser.add_argument("--zeros", type=int, default=0,
                        help="preenche o numero com zeros a esquerda (ex.: --zeros 6)")
    parser.add_argument("--debug", metavar="ARQUIVO",
                        help="lista as tags de um XML e encerra")
    args = parser.parse_args()

    if args.debug:
        modo_debug(Path(args.debug))
        return

    pasta = Path(args.pasta).resolve()
    if not pasta.is_dir():
        sys.exit(f"Pasta nao encontrada: {pasta}")

    saida = pasta if args.mover else (pasta / args.saida)
    if args.aplicar and not args.mover:
        saida.mkdir(exist_ok=True)

    arquivos_xml = sorted(
        p for p in pasta.iterdir()
        if p.suffix.lower() == ".xml" and p.is_file()
    )
    pdfs = sorted(p for p in pasta.iterdir() if p.suffix.lower() == ".pdf")
    if not arquivos_xml:
        sys.exit(f"Nenhum XML em {pasta}")

    linhas, usados, reservados = [], set(), set()
    contagem = {"ok": 0, "sem_dados": 0, "lote": 0, "sem_pdf": 0,
                "ja_existe": 0, "erro": 0}

    for xml in arquivos_xml:
        try:
            raiz = carregar_raiz(xml)
        except ET.ParseError as erro:
            linhas.append({"arquivo_xml": xml.name, "status": f"XML invalido: {erro}"})
            contagem["erro"] += 1
            continue

        layout = detectar_layout(raiz) or LAYOUT_GENERICO
        notas = encontrar_notas(raiz, layout)
        if not notas:
            linhas.append({"arquivo_xml": xml.name, "status": "Nenhuma nota localizada"})
            contagem["erro"] += 1
            continue

        for nota in notas:
            dados = extrair(nota, layout)
            dados["arquivo_xml"] = xml.name
            avisos = []

            if layout is LAYOUT_GENERICO:
                avisos.append("layout nao reconhecido - conferir")
            if not dados.get("numero"):
                avisos.append("numero nao encontrado")
            if not dados.get("prestador_nome"):
                avisos.append("prestador nao encontrado")

            base = montar_nome(dados, args.zeros)

            if len(notas) > 1:
                avisos.append(f"arquivo contem {len(notas)} notas - nao renomeado")
                contagem["lote"] += 1
                dados["status"] = "; ".join(avisos)
                linhas.append(dados)
                continue

            if base is None:
                contagem["sem_dados"] += 1
                dados["status"] = "; ".join(avisos) or "dados insuficientes"
                linhas.append(dados)
                continue

            pdf = achar_pdf(xml, dados, pdfs, usados)
            if pdf is None:
                avisos.append("PDF nao encontrado")
                contagem["sem_pdf"] += 1

            destino_xml, ja_existe = destino(saida, base, ".xml", reservados)
            dados["arquivo_xml"] = destino_xml.name
            destino_pdf = None
            if pdf is not None:
                destino_pdf, _ = destino(saida, base, ".pdf", reservados)
                dados["arquivo_pdf"] = destino_pdf.name
                usados.add(pdf)

            if ja_existe:
                dados["status"] = "ja organizado - ignorado"
                contagem["ja_existe"] += 1
                linhas.append(dados)
                continue

            if args.aplicar:
                acao = shutil.move if args.mover else shutil.copy2
                acao(str(xml), str(destino_xml))
                if pdf is not None:
                    acao(str(pdf), str(destino_pdf))
            else:
                print(f"  {xml.name}  ->  {destino_xml.name}")
                if pdf is not None:
                    print(f"  {pdf.name}  ->  {destino_pdf.name}")

            dados["status"] = "; ".join(avisos) if avisos else "OK"
            if not avisos:
                contagem["ok"] += 1
            linhas.append(dados)

    caminho_planilha = pasta / args.planilha
    gerar_planilha(linhas, caminho_planilha)

    print()
    print(f"Notas processadas ....... {len(linhas)}")
    print(f"  OK .................... {contagem['ok']}")
    print(f"  Sem PDF pareado ....... {contagem['sem_pdf']}")
    print(f"  Em arquivo de lote .... {contagem['lote']}")
    print(f"  Ja organizadas ........ {contagem['ja_existe']}")
    print(f"  Dados insuficientes ... {contagem['sem_dados']}")
    print(f"  Erro de leitura ....... {contagem['erro']}")
    print(f"Planilha ................ {caminho_planilha}")
    if not args.aplicar:
        print("\nSIMULACAO - nenhum arquivo foi alterado. Use --aplicar para executar.")


if __name__ == "__main__":
    main()